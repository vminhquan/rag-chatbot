"""Tạo báo cáo thực tập Word dựa trên mẫu SAD của Đại học Đại Nam."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/Users/vominhquan/Downloads/Mau bao cao BTL - SAD.docx")
OUTPUT = ROOT / "Bao_cao_thuc_tap_AI_Chatbot_RAG.docx"


def clear_document(doc: Document) -> None:
    """Xóa nội dung mẫu nhưng giữ section, header, footer và style gốc."""

    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph) -> None:
    """Thêm trường PAGE động vào footer."""

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_toc(paragraph) -> None:
    """Thêm mục lục tự động; Word cập nhật khi mở và nhấn F9."""

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Nhấn chuột phải và chọn Update Field để cập nhật mục lục."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(4)

    for name, size, color in (
        ("Heading 1", 16, "000000"),
        ("Heading 2", 14, "000000"),
        ("Heading 3", 13, "000000"),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
        footer = section.footer
        if not footer.paragraphs:
            footer.add_paragraph()
        footer.paragraphs[0].clear()
        add_page_number(footer.paragraphs[0])


def centered(doc, text: str, size=14, bold=False, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    return p


def paragraph(doc, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        p.add_run(text[len(bold_prefix) :])
    else:
        p.add_run(text)
    return p


def bullets(doc, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Paragraph")
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.75)
        p.add_run("• " + item)


def table(doc, headers: list[str], rows: list[list[str]], widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], header, True)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in t.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def caption(doc, text: str):
    p = centered(doc, text, size=11, after=4)
    p.runs[0].italic = True


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_cover(doc: Document, grading=False):
    centered(doc, "BỘ GIÁO DỤC VÀ ĐÀO TẠO", 15, True)
    centered(doc, "TRƯỜNG ĐẠI HỌC ĐẠI NAM", 15, True)
    try:
        with ZipFile(TEMPLATE) as archive:
            logo = BytesIO(archive.read("word/media/image1.png"))
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo, width=Inches(1.25))
    except Exception:
        centered(doc, "", 8)
    centered(doc, "BÁO CÁO THỰC TẬP", 20, True, before=12, after=8)
    centered(doc, "ĐỀ TÀI", 14, True)
    centered(doc, "XÂY DỰNG AI CHATBOT TUYỂN DỤNG", 18, True, before=4)
    centered(doc, "ỨNG DỤNG RAG, PINECONE VÀ LLAMA.CPP", 18, True)
    if grading:
        centered(doc, "CÁN BỘ CHẤM BÁO CÁO", 13, True, before=35)
        centered(doc, "1. ........................................................................", 13)
        centered(doc, "2. ........................................................................", 13)
    else:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(4)
        p.paragraph_format.space_before = Pt(24)
        for line in (
            "Sinh viên thực hiện: [HỌ VÀ TÊN]",
            "Mã sinh viên: [MÃ SINH VIÊN]",
            "Lớp: [TÊN LỚP]",
            "Đơn vị thực tập: [TÊN ĐƠN VỊ]",
            "Giảng viên hướng dẫn: [HỌ TÊN GIẢNG VIÊN]",
        ):
            r = p.add_run(line + "\n")
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
    centered(doc, "Hà Nội, năm 2026", 13, before=45)


def add_report(doc: Document):
    # Trang 1–2: bìa theo mẫu.
    add_cover(doc)
    page_break(doc)
    add_cover(doc, grading=True)
    page_break(doc)

    # Trang 3: lời nói đầu.
    doc.add_heading("LỜI NÓI ĐẦU", level=1)
    paragraph(doc, "Trong bối cảnh chuyển đổi số diễn ra mạnh mẽ, trí tuệ nhân tạo đang được ứng dụng vào nhiều hoạt động quản trị doanh nghiệp, trong đó có tuyển dụng nhân sự. Ứng viên thường phải tìm kiếm thông tin phân tán trên website, mạng xã hội và các tin đăng khác nhau để biết doanh nghiệp đang tuyển vị trí nào, yêu cầu công việc ra sao và có những chính sách phúc lợi gì. Quá trình này tốn thời gian, trong khi bộ phận tuyển dụng phải lặp lại nhiều câu trả lời tương tự.")
    paragraph(doc, "Xuất phát từ nhu cầu đó, đề tài xây dựng AI Chatbot tuyển dụng ứng dụng Retrieval-Augmented Generation (RAG) được lựa chọn trong thời gian thực tập. Hệ thống kết hợp cơ sở dữ liệu vector Pinecone, mô hình embedding llama-text-embed-v2 và mô hình ngôn ngữ Qwen2.5-1.5B chạy cục bộ qua llama.cpp. Giao diện ReactJS giao tiếp với REST API FastAPI, cho phép ứng viên chọn công ty và hỏi đáp bằng tiếng Việt dựa trên dữ liệu tuyển dụng đã được kiểm soát.")
    paragraph(doc, "Báo cáo tập trung trình bày cơ sở lý thuyết, quá trình phân tích yêu cầu, kiến trúc hệ thống, cách tổ chức dữ liệu đa công ty, quy trình ingest và truy xuất, thiết kế giao diện, kiểm thử và đánh giá. Các công ty trong dữ liệu minh họa là doanh nghiệp hư cấu nhằm tránh gán thông tin tuyển dụng chưa được xác minh cho tổ chức thực tế.")
    paragraph(doc, "Em xin cảm ơn giảng viên hướng dẫn, đơn vị thực tập và các anh chị đã hỗ trợ kiến thức, góp ý về nghiệp vụ cũng như tạo điều kiện để em hoàn thành đề tài. Do thời gian và nguồn lực còn hạn chế, báo cáo khó tránh khỏi thiếu sót. Em mong nhận được góp ý để tiếp tục hoàn thiện sản phẩm.")
    page_break(doc)

    # Trang 4: mục lục.
    doc.add_heading("MỤC LỤC", level=1)
    add_toc(doc.add_paragraph())
    paragraph(doc, "Lưu ý: mở file trong Microsoft Word, nhấn Ctrl+A rồi F9 để cập nhật mục lục và số trang tự động.")
    page_break(doc)

    # Trang 5: danh mục.
    doc.add_heading("DANH MỤC HÌNH ẢNH", level=1)
    bullets(doc, [
        "Hình 2.1. Quy trình Retrieval-Augmented Generation",
        "Hình 3.1. Kiến trúc tổng thể hệ thống",
        "Hình 3.2. Luồng sequence xử lý một câu hỏi",
        "Hình 3.3. Mô hình phân vùng dữ liệu đa công ty",
    ])
    doc.add_heading("DANH MỤC BẢNG BIỂU", level=1)
    bullets(doc, [
        "Bảng 1.1. Mục tiêu chức năng của đề tài",
        "Bảng 2.1. Công nghệ và công cụ sử dụng",
        "Bảng 3.1. Yêu cầu chức năng", "Bảng 3.2. API của hệ thống",
        "Bảng 3.3. Metadata Pinecone", "Bảng 3.4. Kết quả kiểm thử",
    ])
    doc.add_heading("DANH MỤC CÁC TỪ VIẾT TẮT", level=1)
    table(doc, ["Từ viết tắt", "Giải nghĩa"], [
        ["AI", "Artificial Intelligence – Trí tuệ nhân tạo"],
        ["LLM", "Large Language Model – Mô hình ngôn ngữ lớn"],
        ["RAG", "Retrieval-Augmented Generation"],
        ["API", "Application Programming Interface"],
        ["UI/UX", "User Interface/User Experience"],
        ["JD", "Job Description – Mô tả công việc"],
        ["GGUF", "Định dạng mô hình tối ưu cho llama.cpp"],
    ])
    page_break(doc)

    # Trang 6.
    doc.add_heading("Chương 1. TỔNG QUAN VỀ ĐỀ TÀI", level=1)
    doc.add_heading("1.1. Bối cảnh và lý do lựa chọn", level=2)
    paragraph(doc, "Tuyển dụng là điểm tiếp xúc đầu tiên giữa doanh nghiệp và ứng viên. Chất lượng thông tin ở giai đoạn này ảnh hưởng trực tiếp đến trải nghiệm ứng viên và khả năng thu hút nhân sự phù hợp. Tuy nhiên, thông tin tuyển dụng thường nằm rải rác trong nhiều tài liệu; cách trình bày không đồng nhất và nhanh lỗi thời khi vị trí được mở hoặc đóng.")
    paragraph(doc, "Chatbot dựa hoàn toàn vào LLM có thể giao tiếp tự nhiên nhưng dễ tạo ra thông tin không tồn tại. Đối với tuyển dụng, việc bịa mức lương, phúc lợi hoặc yêu cầu kinh nghiệm gây ảnh hưởng đến uy tín doanh nghiệp. RAG giải quyết vấn đề bằng cách truy xuất tài liệu có liên quan trước, sau đó yêu cầu mô hình chỉ trả lời từ ngữ cảnh tìm được.")
    doc.add_heading("1.2. Mục tiêu đề tài", level=2)
    table(doc, ["Mục tiêu", "Kết quả mong đợi"], [
        ["Hỏi đáp tuyển dụng", "Trả lời vị trí, JD, yêu cầu, lương và phúc lợi"],
        ["Kiểm soát tính đúng", "Không trả lời ngoài dữ liệu đã ingest"],
        ["Đa công ty", "Cách ly dữ liệu bằng namespace riêng"],
        ["Chi phí LLM", "Dùng llama.cpp chạy mô hình local, không dùng quota OpenAI"],
        ["Khả năng tích hợp", "REST API cho frontend ReactJS"],
    ])
    caption(doc, "Bảng 1.1. Mục tiêu chức năng của đề tài")
    page_break(doc)

    # Trang 7.
    doc.add_heading("1.3. Phạm vi và đối tượng sử dụng", level=2)
    paragraph(doc, "Đối tượng sử dụng chính là ứng viên muốn tìm hiểu doanh nghiệp trước khi nộp hồ sơ. Người dùng chọn một công ty trong danh mục, sau đó đặt câu hỏi tự nhiên bằng tiếng Việt. Phạm vi dữ liệu gồm giới thiệu doanh nghiệp, địa điểm, văn hóa, phúc lợi, vị trí đang tuyển hoặc đã đóng, mô tả công việc, yêu cầu, cấp bậc và mức lương nếu được công bố.")
    paragraph(doc, "Đối tượng quản trị là nhà tuyển dụng hoặc người vận hành hệ thống. Trong phiên bản hiện tại, quản trị dữ liệu được thực hiện thông qua tệp companies.json và script dòng lệnh. Các chức năng quản trị trên giao diện, xác thực tài khoản, nhận CV, đặt lịch phỏng vấn và kết nối ATS chưa thuộc phạm vi triển khai.")
    doc.add_heading("1.4. Phương pháp thực hiện", level=2)
    bullets(doc, [
        "Khảo sát bài toán và xác định nhóm câu hỏi thường gặp của ứng viên.",
        "Nghiên cứu kiến trúc RAG, vector database và mô hình ngôn ngữ chạy local.",
        "Thiết kế API, schema dữ liệu và mô hình namespace đa công ty.",
        "Xây dựng backend FastAPI, pipeline Pinecone và llama.cpp.",
        "Phát triển giao diện ReactJS theo hướng hội thoại dẫn dắt.",
        "Kiểm thử từng tầng và đánh giá kết quả trả lời thực tế.",
    ])
    doc.add_heading("1.5. Kế hoạch thực tập", level=2)
    table(doc, ["Giai đoạn", "Nội dung"], [
        ["Tuần 1–2", "Khảo sát nghiệp vụ, học RAG và Pinecone"],
        ["Tuần 3–4", "Thiết kế dữ liệu, xây ingest và retrieval"],
        ["Tuần 5–6", "Xây FastAPI, tích hợp llama.cpp"],
        ["Tuần 7–8", "Xây ReactJS, kiểm thử và viết báo cáo"],
    ])
    page_break(doc)

    # Trang 8.
    doc.add_heading("1.6. Kết quả dự kiến và tiêu chí đánh giá", level=2)
    paragraph(doc, "Sản phẩm phải vận hành theo chuỗi hoàn chỉnh từ dữ liệu nguồn đến câu trả lời trên giao diện. Một câu trả lời được xem là đạt khi thông tin tồn tại trong catalog, đúng công ty đã chọn, không trộn dữ liệu giữa các namespace và có thể truy vết tới nguồn. Với câu hỏi dạng danh sách, hệ thống phải liệt kê đầy đủ tất cả vị trí có trạng thái open.")
    paragraph(doc, "Các tiêu chí kỹ thuật gồm API phản hồi đúng schema, frontend không lưu khóa bí mật, Pinecone tự tạo embedding khi upsert và search, llama-server chỉ bind localhost, giao diện hỗ trợ Markdown và trạng thái đang xử lý. Hệ thống cần hoạt động trên môi trường phát triển VS Code và có các lệnh chạy độc lập cho vector database, LLM server, backend và frontend.")
    doc.add_heading("1.7. Ý nghĩa thực tiễn", level=2)
    paragraph(doc, "Giải pháp giúp ứng viên tiếp cận thông tin tuyển dụng nhanh hơn và giảm số câu hỏi lặp lại cho bộ phận nhân sự. Kiến trúc đa tenant cho phép một nền tảng phục vụ nhiều công ty mà vẫn giữ được cách ly dữ liệu. Việc chạy LLM tại máy cục bộ phù hợp với bài toán học tập, thử nghiệm nội bộ và những trường hợp cần kiểm soát dữ liệu đầu vào.")
    paragraph(doc, "Đề tài cũng là cơ hội thực hành tổng hợp nhiều kỹ năng: phân tích yêu cầu, thiết kế kiến trúc, phát triển REST API, xử lý dữ liệu phi cấu trúc, tích hợp dịch vụ AI, quản lý trạng thái frontend và kiểm thử end-to-end.")
    page_break(doc)

    # Trang 9.
    doc.add_heading("Chương 2. CƠ SỞ LÝ THUYẾT", level=1)
    doc.add_heading("2.1. Mô hình ngôn ngữ lớn và chatbot", level=2)
    paragraph(doc, "LLM là mô hình học máy được huấn luyện trên lượng lớn văn bản để dự đoán và sinh chuỗi ngôn ngữ. Trong chatbot, LLM tiếp nhận instruction, lịch sử hội thoại và câu hỏi để sinh phản hồi. Khả năng diễn đạt tốt không đồng nghĩa với dữ liệu luôn chính xác; mô hình có thể tạo câu trả lời hợp lý về mặt ngôn ngữ nhưng sai về mặt sự kiện.")
    paragraph(doc, "Đề tài sử dụng Qwen2.5-1.5B-Instruct định dạng GGUF, lượng tử hóa Q4_K_M. Mô hình được phục vụ bằng llama.cpp qua endpoint tương thích OpenAI Chat Completions. Cấu hình alias qwen2.5-1.5b-instruct, context 4096 và localhost:8080 giúp backend gọi mô hình như một dịch vụ HTTP cục bộ.")
    doc.add_heading("2.2. Retrieval-Augmented Generation", level=2)
    paragraph(doc, "RAG kết hợp truy xuất thông tin với sinh ngôn ngữ. Tài liệu được chia thành record, chuyển thành vector và lưu trong vector database. Khi có câu hỏi, hệ thống tạo vector truy vấn, tìm các record gần nhất rồi ghép chúng vào prompt. Mô hình sinh câu trả lời dựa trên tập ngữ cảnh này thay vì chỉ dựa vào tri thức đã học.")
    table(doc, ["1. Dữ liệu", "2. Truy xuất", "3. Sinh câu trả lời"], [["JD/phúc lợi → embedding", "Câu hỏi → top-k records", "Context + lịch sử → llama.cpp"]])
    caption(doc, "Hình 2.1. Quy trình Retrieval-Augmented Generation")
    page_break(doc)

    # Trang 10.
    doc.add_heading("2.3. Embedding và tìm kiếm ngữ nghĩa", level=2)
    paragraph(doc, "Embedding biểu diễn văn bản dưới dạng vector số sao cho nội dung gần nghĩa có khoảng cách nhỏ. Khác với tìm kiếm từ khóa, semantic search có thể nhận ra quan hệ giữa câu hỏi và JD dù người dùng không sử dụng đúng cụm từ trong tài liệu. Index của đề tài gắn trực tiếp model llama-text-embed-v2 của Pinecone; trường text được embedding tự động khi upsert và khi search.")
    doc.add_heading("2.4. Pinecone và namespace", level=2)
    paragraph(doc, "Pinecone là vector database được quản lý, cung cấp index serverless, metadata filter và integrated inference. Mỗi công ty được ánh xạ thành một namespace company-{company_id}; JD công khai đồng thời được lưu ở public-jobs. Cách tổ chức này giảm nguy cơ truy xuất nhầm dữ liệu và cho phép mở rộng nhiều tenant trên cùng index.")
    doc.add_heading("2.5. llama.cpp", level=2)
    paragraph(doc, "llama.cpp là runtime C/C++ tối ưu để chạy mô hình GGUF trên máy cá nhân. llama-server cung cấp HTTP API, quản lý context và thực hiện suy luận mà không gửi prompt tới nhà cung cấp LLM bên ngoài. Nhược điểm là chất lượng và tốc độ phụ thuộc model, cấu hình lượng tử hóa và phần cứng. Qwen 1.5B có ưu thế nhẹ nhưng có thể bỏ sót mục khi tổng hợp danh sách dài.")
    doc.add_heading("2.6. FastAPI và ReactJS", level=2)
    paragraph(doc, "FastAPI cung cấp validation bằng Pydantic, tài liệu Swagger và middleware CORS. ReactJS quản lý lịch sử hội thoại, trạng thái loading/error/success, lựa chọn công ty và render Markdown. Frontend chỉ gọi FastAPI; khóa Pinecone không xuất hiện trong trình duyệt.")
    page_break(doc)

    # Trang 11.
    doc.add_heading("2.7. Công nghệ và công cụ sử dụng", level=2)
    table(doc, ["Công nghệ", "Vai trò", "Lý do lựa chọn"], [
        ["Python 3.10", "Ngôn ngữ backend", "Hệ sinh thái AI tốt"],
        ["FastAPI", "REST API", "Nhanh, schema rõ, có Swagger"],
        ["Pinecone", "Vector database", "Integrated embedding và namespace"],
        ["llama-text-embed-v2", "Embedding", "Tối ưu retrieval, Pinecone quản lý"],
        ["llama.cpp", "LLM runtime", "Chạy local, hỗ trợ GGUF"],
        ["Qwen2.5-1.5B", "Sinh câu trả lời", "Nhẹ, hỗ trợ tiếng Việt ở mức thử nghiệm"],
        ["React + Vite", "Giao diện", "Phản hồi nhanh, build đơn giản"],
        ["VS Code", "Môi trường phát triển", "Quản lý source và terminal thuận tiện"],
    ])
    caption(doc, "Bảng 2.1. Công nghệ và công cụ sử dụng")
    doc.add_heading("2.8. An toàn và tính đúng của câu trả lời", level=2)
    paragraph(doc, "Prompt hệ thống yêu cầu chỉ trả lời từ ngữ cảnh, không tự tạo mức lương hoặc phúc lợi và phải thông báo khi thiếu dữ liệu. Backend dùng allow-list company_id để ngăn frontend truy cập namespace tùy ý. Riêng intent liệt kê vị trí đang tuyển được xử lý có tính quyết định từ catalog nhằm khắc phục giới hạn bỏ sót của model nhỏ.")
    page_break(doc)

    # Trang 12.
    doc.add_heading("Chương 3. PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG", level=1)
    doc.add_heading("3.1. Phân tích đặc tả yêu cầu", level=2)
    table(doc, ["Mã", "Yêu cầu", "Tác nhân"], [
        ["FR-01", "Hiển thị danh sách công ty", "Ứng viên"],
        ["FR-02", "Chọn công ty trước khi trò chuyện", "Ứng viên"],
        ["FR-03", "Hỏi vị trí, JD, yêu cầu, lương, phúc lợi", "Ứng viên"],
        ["FR-04", "Hiển thị nguồn tham khảo", "Ứng viên"],
        ["FR-05", "Reset hội thoại", "Ứng viên"],
        ["FR-06", "Ingest tất cả hoặc một công ty", "Quản trị"],
        ["FR-07", "Cách ly dữ liệu theo namespace", "Hệ thống"],
    ])
    caption(doc, "Bảng 3.1. Yêu cầu chức năng")
    doc.add_heading("3.1.1. Yêu cầu phi chức năng", level=3)
    bullets(doc, [
        "Tính đúng: câu trả lời không được vượt ngoài dữ liệu nguồn.",
        "Bảo mật: khóa API chỉ tồn tại phía backend; llama-server bind localhost.",
        "Khả dụng: giao diện responsive, có trạng thái typing và thông báo lỗi.",
        "Mở rộng: thêm công ty qua catalog và namespace mà không tạo index mới.",
        "Bảo trì: tách config, schema, Pinecone service, RAG service và script ingest.",
    ])
    page_break(doc)

    # Trang 13.
    doc.add_heading("3.2. Kiến trúc tổng thể", level=2)
    table(doc, ["Tầng giao diện", "Tầng dịch vụ", "Tầng AI và dữ liệu"], [
        ["ReactJS/Vite\nCompany picker\nMarkdown chat", "FastAPI\nSchema validation\nRAG orchestration", "Pinecone\nllama-text-embed-v2\nllama.cpp + Qwen"]
    ])
    caption(doc, "Hình 3.1. Kiến trúc tổng thể hệ thống")
    paragraph(doc, "Frontend tải catalog từ GET /api/catalog và yêu cầu ứng viên chọn công ty. Khi gửi câu hỏi, React chuyển question, company_id và history qua POST /api/chat. Backend xác thực company_id, chuyển thành namespace, truy xuất top-k record từ Pinecone và xây context. Nếu câu hỏi là danh sách vị trí đang tuyển, backend trả trực tiếp dữ liệu cấu trúc. Các intent còn lại được gửi tới llama-server để sinh câu trả lời.")
    doc.add_heading("3.2.1. Các API chính", level=3)
    table(doc, ["Method", "Đường dẫn", "Chức năng"], [
        ["GET", "/api/catalog", "Danh sách công ty/ngành/địa điểm"],
        ["POST", "/api/chat", "Xử lý câu hỏi RAG"],
        ["GET", "llama.cpp /health", "Kiểm tra LLM server"],
        ["POST", "llama.cpp /v1/chat/completions", "Sinh phản hồi local"],
    ])
    caption(doc, "Bảng 3.2. API của hệ thống")
    page_break(doc)

    # Trang 14.
    doc.add_heading("3.2.2. Luồng xử lý câu hỏi", level=3)
    table(doc, ["Bước", "Thành phần", "Xử lý"], [
        ["1", "ReactJS", "Gửi question, company_id, history"],
        ["2", "FastAPI", "Validate schema và resolve namespace"],
        ["3", "Pinecone", "Embedding câu hỏi và semantic search"],
        ["4", "RAG service", "Chuẩn hóa nguồn, xây context"],
        ["5", "Intent router", "Danh sách việc làm → catalog; còn lại → LLM"],
        ["6", "llama.cpp", "Sinh câu trả lời từ context"],
        ["7", "ReactJS", "Render Markdown, thời gian và nguồn"],
    ])
    caption(doc, "Hình 3.2. Luồng sequence xử lý một câu hỏi")
    doc.add_heading("3.3. Thiết kế dữ liệu đa công ty", level=2)
    paragraph(doc, "Catalog companies.json lưu dữ liệu nguồn có cấu trúc. Mỗi công ty gồm id, tên, ngành, địa điểm, giới thiệu, phúc lợi và danh sách jobs. Mỗi job có mã, tiêu đề, nhóm nghề, cấp bậc, địa điểm, loại hình, trạng thái, mô tả, yêu cầu và lương.")
    table(doc, ["Index recruitment-chatbot", "Nội dung"], [
        ["company-novatech", "Giới thiệu, phúc lợi, JD NovaTech"],
        ["company-antam-health", "Giới thiệu, phúc lợi, JD An Tâm"],
        ["company-greenroute", "Giới thiệu, phúc lợi, JD GreenRoute"],
        ["company-mocviet-retail", "Giới thiệu, phúc lợi, JD Mộc Việt"],
        ["public-jobs", "Bản tổng hợp JD để tìm xuyên công ty"],
    ])
    caption(doc, "Hình 3.3. Mô hình phân vùng dữ liệu đa công ty")
    page_break(doc)

    # Trang 15.
    doc.add_heading("3.3.1. Metadata record", level=3)
    table(doc, ["Trường", "Ý nghĩa", "Ví dụ"], [
        ["_id", "ID ổn định SHA-256", "f82a..."],
        ["text", "Nội dung được embedding", "VỊ TRÍ: Backend Developer..."],
        ["company_id", "Mã công ty", "novatech"],
        ["job_title", "Tên vị trí", "Backend Developer"],
        ["industry", "Ngành", "Công nghệ thông tin"],
        ["location", "Địa điểm", "TP. Hồ Chí Minh"],
        ["status", "Trạng thái", "open"],
        ["source", "Nguồn truy vết", "companies.json#..."],
    ])
    caption(doc, "Bảng 3.3. Metadata Pinecone")
    doc.add_heading("3.4. Thiết kế pipeline ingest", level=2)
    paragraph(doc, "Script ingest_data.py đọc catalog, tạo hai record company/benefits và một record cho mỗi job. ID được băm ổn định để chạy lại sẽ cập nhật thay vì tạo bản sao. Record đầy đủ được upsert vào namespace công ty; riêng record_type=job được sao chép sang public-jobs. Pinecone tự embedding trường text bằng model gắn với index.")
    paragraph(doc, "Script hỗ trợ ingest toàn bộ bằng python -m scripts.ingest_data hoặc một công ty bằng tham số --company. Trong lần triển khai thử nghiệm, bốn công ty tạo 16 record riêng và tám JD tổng hợp. Dữ liệu mới có thể cần vài giây trước khi xuất hiện do cơ chế eventual consistency.")
    doc.add_heading("3.5. Xử lý RAG và prompt", level=2)
    paragraph(doc, "RAG service giới hạn lịch sử ở sáu tin nhắn gần nhất, định dạng từng tài liệu với tên công ty, vị trí và địa điểm. System instruction yêu cầu mô hình đọc toàn bộ context, phân biệt open/closed và từ chối khi thiếu thông tin. Nhiệt độ 0,2 giúp câu trả lời ổn định hơn.")
    page_break(doc)

    # Trang 16.
    doc.add_heading("3.6. Thiết kế giao diện", level=2)
    paragraph(doc, "Giao diện được xây bằng React và Vite theo thiết kế modern-minimal. Khi mở ứng dụng, bot hỏi người dùng muốn tìm hiểu công ty nào. Danh sách công ty hiển thị ngay trong message area; ô nhập bị khóa cho tới khi có lựa chọn. Sau khi chọn, bot xác nhận và cung cấp ba câu hỏi nhanh: vị trí đang tuyển, giới thiệu công ty và phúc lợi.")
    paragraph(doc, "Tin nhắn người dùng nằm bên phải với màu chủ đạo; tin nhắn bot nằm bên trái, hỗ trợ Markdown, danh sách và code block. Mỗi tin có thời gian; câu trả lời có khu vực mở xem nguồn. Header hiển thị trạng thái online/typing và nút reset. Textarea tự tăng chiều cao; Enter gửi, Shift+Enter xuống dòng. Giao diện chuyển toàn màn hình trên thiết bị nhỏ.")
    doc.add_heading("3.7. Giải pháp khắc phục bỏ sót danh sách", level=2)
    paragraph(doc, "Trong kiểm thử ban đầu, câu hỏi 'Đang tuyển vị trí nào?' tại NovaTech được Pinecone trả đủ Backend và Frontend nhưng Qwen 1.5B chỉ nêu Frontend. Câu hỏi lương tiếp theo lại đọc được cả hai vị trí, chứng minh dữ liệu không thiếu mà lỗi nằm ở bước tổng hợp của model nhỏ.")
    paragraph(doc, "Giải pháp là tách intent dữ liệu cấu trúc. Hàm is_open_job_list_question nhận diện các mẫu hỏi liệt kê, sau đó build_open_job_list_answer đọc trực tiếp catalog và lọc status=open. Cách làm deterministic bảo đảm không bỏ sót; Pinecone vẫn được dùng để cung cấp source. Kết quả sau sửa trả đúng hai vị trí Backend Developer và Frontend Developer.")
    doc.add_heading("3.8. Bảo mật", level=2)
    bullets(doc, [
        "PINECONE_API_KEY chỉ lưu trong .env phía backend.",
        "Frontend không gọi trực tiếp Pinecone hoặc llama-server.",
        "company_id được kiểm tra bằng allow-list từ catalog.",
        "llama-server chỉ lắng nghe 127.0.0.1:8080.",
        "Exception nội bộ không được trả nguyên văn cho trình duyệt production.",
    ])
    page_break(doc)

    # Trang 17.
    doc.add_heading("3.9. KIỂM THỬ VÀ ĐÁNH GIÁ", level=2)
    table(doc, ["TC", "Dữ liệu kiểm thử", "Kết quả mong đợi", "Trạng thái"], [
        ["01", "Chọn NovaTech", "Bot xác nhận đúng công ty", "Đạt"],
        ["02", "Đang tuyển vị trí nào?", "BE và FE, không thiếu mục", "Đạt"],
        ["03", "Mức lương ở đây?", "BE 25–40, FE 22–35 triệu", "Đạt"],
        ["04", "Phúc lợi NovaTech", "Chỉ dùng record benefits", "Đạt"],
        ["05", "Supply Chain tại Đà Nẵng", "GreenRoute – Supply Chain Analyst", "Đạt"],
        ["06", "Công ty không tồn tại", "HTTP 400", "Đạt"],
        ["07", "Dừng llama-server", "Backend báo lỗi có kiểm soát", "Đạt"],
        ["08", "Màn hình 375px", "Không cuộn ngang", "Đạt"],
    ])
    caption(doc, "Bảng 3.4. Kết quả kiểm thử chức năng")
    doc.add_heading("3.9.1. Kiểm thử truy xuất", level=3)
    paragraph(doc, "Truy vấn thực tế 'Có việc phân tích chuỗi cung ứng nào?' với namespace public-jobs, location=Đà Nẵng và status=open trả một kết quả GreenRoute Logistics – Supply Chain Analyst, score 0,2651. Truy vấn danh sách NovaTech trả hai source open tương ứng Frontend và Backend.")
    doc.add_heading("3.9.2. Kiểm thử build", level=3)
    paragraph(doc, "Backend được kiểm tra cú pháp và import bằng Python 3.10. Frontend production build bằng Vite thành công với 2.053 module được transform. API schema mặc định company_id=all nhưng giao diện hiện yêu cầu chọn một công ty trước khi gửi câu hỏi.")
    doc.add_heading("3.9.3. Đánh giá", level=3)
    paragraph(doc, "Hệ thống đáp ứng luồng MVP, dữ liệu được cách ly và câu hỏi phổ biến có kết quả đúng. Hạn chế chính là tốc độ suy luận local khoảng vài token mỗi giây trên bản build thử nghiệm và năng lực tổng hợp của model 1.5B. Similarity score cũng không nên dùng như xác suất đúng tuyệt đối mà cần đánh giá trên bộ câu hỏi thực tế.")
    page_break(doc)

    # Trang 18.
    doc.add_heading("KẾT LUẬN", level=1)
    doc.add_heading("1. Kết quả đạt được", level=2)
    paragraph(doc, "Đề tài đã xây dựng thành công chatbot tuyển dụng RAG có giao diện ReactJS, REST API FastAPI, Pinecone integrated embedding và LLM Qwen chạy local qua llama.cpp. Hệ thống hỗ trợ nhiều công ty, nhiều nhóm nghề, hiển thị nguồn, quản lý lịch sử ngắn và bảo đảm danh sách vị trí open bằng xử lý deterministic.")
    doc.add_heading("2. Nhược điểm", level=2)
    paragraph(doc, "Dữ liệu hiện được quản trị qua JSON và chưa có trang quản trị. Model nhỏ có chất lượng thấp hơn dịch vụ LLM thương mại, chưa có streaming token và chưa đo lường tự động precision/recall của retrieval. Chức năng đính kèm trên frontend chưa có API nhận CV; hệ thống chưa triển khai xác thực, phân quyền và audit log.")
    doc.add_heading("3. Hướng phát triển", level=2)
    bullets(doc, [
        "Xây trang quản trị để thêm công ty, JD và đóng/mở vị trí.",
        "Bổ sung parser PDF/DOCX và chunk theo cấu trúc tiêu đề.",
        "Dùng model local lớn hơn hoặc reranker để nâng chất lượng.",
        "Streaming câu trả lời, lưu conversation và thu thập phản hồi.",
        "Tích hợp ATS, nhận CV có kiểm soát và đặt lịch phỏng vấn.",
        "Xây bộ đánh giá RAG định kỳ và giám sát latency/error rate.",
    ])
    doc.add_heading("TÀI LIỆU THAM KHẢO", level=1)
    refs = [
        "[1] Pinecone Documentation, Integrated embedding and multitenancy, https://docs.pinecone.io/.",
        "[2] llama.cpp, LLM inference in C/C++, https://github.com/ggml-org/llama.cpp.",
        "[3] FastAPI Documentation, https://fastapi.tiangolo.com/.",
        "[4] React Documentation, https://react.dev/.",
        "[5] Lewis et al., Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks, 2020.",
        "[6] Qwen Team, Qwen2.5 Technical Report and model documentation.",
    ]
    for ref in refs:
        p = doc.add_paragraph(ref)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.15


def main():
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Không tìm thấy file mẫu: {TEMPLATE}")
    doc = Document(TEMPLATE)
    clear_document(doc)
    configure_styles(doc)
    add_report(doc)
    # Yêu cầu Word cập nhật field TOC/page khi mở file.
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings._element.append(update_fields)
    doc.core_properties.title = "Báo cáo thực tập – AI Chatbot tuyển dụng RAG"
    doc.core_properties.subject = "ReactJS, FastAPI, Pinecone và llama.cpp"
    doc.core_properties.author = "[HỌ VÀ TÊN SINH VIÊN]"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
